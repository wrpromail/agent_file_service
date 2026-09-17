"""Spreadsheet → structured JSON tables (headers + rows) for agents.

Profile: ``tables``. Deterministic openpyxl/xlrd extract — not markdown reverse-parse.
Keeps max_cols bounded so dirty used-range (max_column=16384) cannot explode memory.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from app.parsers.base import ParseContext, Parser
from app.schemas.content import (
    ContentBlock,
    ContentSource,
    ContentTable,
    ContentWarning,
    DocumentInfo,
    StructuredContent,
)

SCHEMA_VERSION = "spreadsheet_tables@v1"
DEFAULT_MAX_COLS = 64
DEFAULT_MAX_ROWS = 5000
DEFAULT_MAX_SHEETS = 50


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value != value:  # NaN
        return None
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text if text else None


def _trim_trailing_none(cells: list[str | None]) -> list[str | None]:
    while cells and cells[-1] is None:
        cells.pop()
    return cells


def extract_workbook_tables(
    file_path: str,
    *,
    max_sheets: int = DEFAULT_MAX_SHEETS,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLS,
) -> dict[str, Any]:
    """Return spreadsheet_tables@v1 dict from an on-disk workbook path."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    notes: list[str] = []
    if suffix == ".xls":
        sheets, extract_notes = _extract_xls(
            path, max_sheets=max_sheets, max_rows=max_rows, max_cols=max_cols
        )
    else:
        sheets, extract_notes = _extract_xlsx(
            path, max_sheets=max_sheets, max_rows=max_rows, max_cols=max_cols
        )
    notes.extend(extract_notes)
    return {
        "schema_version": SCHEMA_VERSION,
        "filename": path.name,
        "sheet_count": len(sheets),
        "limits": {
            "max_sheets": max_sheets,
            "max_rows_per_sheet": max_rows,
            "max_cols": max_cols,
        },
        "sheets": sheets,
        "notes": notes,
    }


def _extract_xlsx(
    path: Path,
    *,
    max_sheets: int,
    max_rows: int,
    max_cols: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "openpyxl is required for tables profile; install openpyxl"
        ) from exc

    notes: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if len(names) > max_sheets:
            notes.append(f"sheet_count {len(names)} > max_sheets {max_sheets}; truncated")
        names = names[:max_sheets]
        out: list[dict[str, Any]] = []
        for name in names:
            ws = wb[name]
            sheet_notes: list[str] = []
            headers: list[str | None] = []
            for row in ws.iter_rows(
                min_row=1, max_row=1, min_col=1, max_col=max_cols, values_only=True
            ):
                headers = _trim_trailing_none([_cell_str(c) for c in (row or ())])
            if len([h for h in headers if h]) >= max_cols:
                sheet_notes.append(f"header truncated to max_cols={max_cols}")
            width = max(len(headers), 1)
            width = min(width, max_cols)

            rows: list[list[str | None]] = []
            truncated = False
            scanned = 0
            for row in ws.iter_rows(
                min_row=2,
                max_row=1 + max_rows + 1,  # allow detecting overflow
                min_col=1,
                max_col=width,
                values_only=True,
            ):
                scanned += 1
                if scanned > max_rows:
                    truncated = True
                    sheet_notes.append(f"rows truncated at max_rows_per_sheet={max_rows}")
                    break
                cells = _trim_trailing_none([_cell_str(c) for c in (row or ())[:width]])
                # pad/truncate to header width for stable columns
                if len(cells) < width:
                    cells = cells + [None] * (width - len(cells))
                elif len(cells) > width:
                    cells = cells[:width]
                if any(c is not None for c in cells):
                    rows.append(cells)

            header_labels = [h if h is not None else f"col_{i+1}" for i, h in enumerate(headers)]
            # if no header cells, still export raw first width
            if not any(headers):
                header_labels = [f"col_{i+1}" for i in range(width)]
                sheet_notes.append("empty_header_row; using col_N labels")

            out.append(
                {
                    "name": name,
                    "headers": header_labels,
                    "rows": [[("" if c is None else c) for c in r] for r in rows],
                    "row_count": len(rows),
                    "header_cols": len(header_labels),
                    "truncated": truncated,
                    "notes": sheet_notes,
                }
            )
        return out, notes
    finally:
        wb.close()


def _extract_xls(
    path: Path,
    *,
    max_sheets: int,
    max_rows: int,
    max_cols: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("xlrd is required for .xls tables profile") from exc

    notes: list[str] = []
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        names = list(book.sheet_names())
        if len(names) > max_sheets:
            notes.append(f"sheet_count {len(names)} > max_sheets {max_sheets}; truncated")
        names = names[:max_sheets]
        out: list[dict[str, Any]] = []
        for name in names:
            sheet = book.sheet_by_name(name)
            sheet_notes: list[str] = []
            ncols = min(int(sheet.ncols or 0), max_cols)
            if int(sheet.ncols or 0) > max_cols:
                sheet_notes.append(f"header truncated to max_cols={max_cols}")
            headers: list[str | None] = []
            if sheet.nrows >= 1:
                headers = _trim_trailing_none(
                    [_cell_str(sheet.cell_value(0, c)) for c in range(ncols)]
                )
            width = max(len(headers), 1)
            width = min(width, max_cols)
            header_labels = [h if h is not None else f"col_{i+1}" for i, h in enumerate(headers)]
            if not any(headers):
                header_labels = [f"col_{i+1}" for i in range(width)]
                sheet_notes.append("empty_header_row; using col_N labels")

            rows: list[list[str]] = []
            truncated = False
            data_nrows = max(0, int(sheet.nrows) - 1)
            limit = min(data_nrows, max_rows)
            if data_nrows > max_rows:
                truncated = True
                sheet_notes.append(f"rows truncated at max_rows_per_sheet={max_rows}")
            for r in range(1, 1 + limit):
                cells = [_cell_str(sheet.cell_value(r, c)) for c in range(width)]
                if any(c is not None for c in cells):
                    rows.append([("" if c is None else c) for c in cells])
            out.append(
                {
                    "name": name,
                    "headers": header_labels,
                    "rows": rows,
                    "row_count": len(rows),
                    "header_cols": len(header_labels),
                    "truncated": truncated,
                    "notes": sheet_notes,
                }
            )
            try:
                book.unload_sheet(name)
            except Exception:
                pass
        return out, notes
    finally:
        book.release_resources()


def _slug(text: str, fallback: str) -> str:
    cleaned = re.sub(r"[^\w\-]+", "-", text.strip(), flags=re.UNICODE)
    cleaned = re.sub(r"-+", "-", cleaned).strip("-").lower()
    return cleaned[:64] or fallback


class SpreadsheetTablesParser(Parser):
    """Produce spreadsheet_tables@v1 JSON for agent/code consumption."""

    name = "spreadsheet_tables"
    version = "0.1.0"
    profile = "tables"
    supported_mime_types = [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
        "text/csv",
    ]

    def __init__(
        self,
        *,
        max_sheets: int = DEFAULT_MAX_SHEETS,
        max_rows: int = DEFAULT_MAX_ROWS,
        max_cols: int = DEFAULT_MAX_COLS,
    ) -> None:
        self.max_sheets = max_sheets
        self.max_rows = max_rows
        self.max_cols = max_cols

    def can_handle(self, mime_type: str) -> bool:
        if not mime_type:
            return True
        if mime_type in self.supported_mime_types:
            return True
        return mime_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }

    def parse(self, file_path: str, context: ParseContext) -> StructuredContent:
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            doc = self._csv_tables(path)
        else:
            doc = extract_workbook_tables(
                file_path,
                max_sheets=self.max_sheets,
                max_rows=self.max_rows,
                max_cols=self.max_cols,
            )
        doc["filename"] = context.filename or path.name

        tables: list[ContentTable] = []
        blocks: list[ContentBlock] = []
        for index, sheet in enumerate(doc.get("sheets") or [], start=1):
            name = str(sheet.get("name") or f"sheet{index}")
            headers = list(sheet.get("headers") or [])
            rows = list(sheet.get("rows") or [])
            # ContentTable.rows includes header as first row for consumers that expect matrix
            matrix = [headers] + rows
            tables.append(
                ContentTable(
                    id=f"table-{_slug(name, f'{index:03d}')}",
                    page=index,
                    rows=[[str(c) if c is not None else "" for c in r] for r in matrix],
                    confidence=1.0,
                    metadata={
                        "sheet_name": name,
                        "header_cols": sheet.get("header_cols"),
                        "row_count": sheet.get("row_count"),
                        "truncated": sheet.get("truncated"),
                        "format": "spreadsheet_tables",
                    },
                )
            )
            # Compact block text: headers only (full rows in tables + metadata.tables_json)
            preview_cols = headers[:8]
            blocks.append(
                ContentBlock(
                    id=f"sheet-{_slug(name, f'{index:03d}')}",
                    type="table",
                    page=index,
                    text=f"## {name}\nheaders: {', '.join(str(h) for h in preview_cols)}",
                    confidence=1.0,
                    metadata={
                        "sheet_name": name,
                        "format": "spreadsheet_tables",
                        "row_count": sheet.get("row_count"),
                    },
                )
            )

        warnings: list[ContentWarning] = []
        if not tables:
            warnings.append(
                ContentWarning(
                    block_id=None,
                    type="parse_partial",
                    detail="no sheets/tables extracted",
                )
            )

        return StructuredContent(
            content_hash=context.content_hash,
            source=ContentSource(
                parser_name=self.name,
                parser_version=self.version,
                profile=self.profile,
            ),
            document=DocumentInfo(page_count=len(tables)),
            blocks=blocks,
            tables=tables,
            warnings=warnings,
            metadata={
                "format": "spreadsheet_tables_json",
                "schema_version": SCHEMA_VERSION,
                "tables_json": doc,
                "filename": context.filename,
                "mime_type": context.mime_type,
                "sheet_count": doc.get("sheet_count"),
                "sheet_names": [s.get("name") for s in (doc.get("sheets") or [])],
            },
        )

    def _csv_tables(self, path: Path) -> dict[str, Any]:
        import csv

        notes: list[str] = []
        with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.reader(fh)
            all_rows = list(reader)
        if not all_rows:
            return {
                "schema_version": SCHEMA_VERSION,
                "filename": path.name,
                "sheet_count": 0,
                "limits": {
                    "max_sheets": self.max_sheets,
                    "max_rows_per_sheet": self.max_rows,
                    "max_cols": self.max_cols,
                },
                "sheets": [],
                "notes": ["empty_csv"],
            }
        header = all_rows[0][: self.max_cols]
        data = all_rows[1:]
        truncated = False
        if len(data) > self.max_rows:
            data = data[: self.max_rows]
            truncated = True
            notes.append(f"rows truncated at max_rows_per_sheet={self.max_rows}")
        if any(len(r) > self.max_cols for r in all_rows):
            notes.append(f"cols truncated to max_cols={self.max_cols}")
            data = [r[: self.max_cols] for r in data]
        return {
            "schema_version": SCHEMA_VERSION,
            "filename": path.name,
            "sheet_count": 1,
            "limits": {
                "max_sheets": self.max_sheets,
                "max_rows_per_sheet": self.max_rows,
                "max_cols": self.max_cols,
            },
            "sheets": [
                {
                    "name": path.stem or "Sheet1",
                    "headers": header,
                    "rows": data,
                    "row_count": len(data),
                    "header_cols": len(header),
                    "truncated": truncated,
                    "notes": [],
                }
            ],
            "notes": notes,
        }
