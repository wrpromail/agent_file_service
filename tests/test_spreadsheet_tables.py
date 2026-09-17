"""Unit tests + API tests for profile=tables spreadsheet JSON extract."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from app.parsers.base import ParseContext
from app.parsers.spreadsheet_tables import SpreadsheetTablesParser, extract_workbook_tables
from app.services.parse_service import SUPPORTED_PARSE_PROFILES
from tests.conftest import reload_parsers, upload_bytes


def _write_multi_sheet_xlsx(path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    apps = wb.active
    apps.title = "应用系统"
    apps.append(["名称(*)", "类型"])
    apps.append(["系统A", "业务"])
    apps.append(["系统B", "支撑"])

    servers = wb.create_sheet("服务器")
    servers.append(["主机名", "IP"])
    for i in range(10):
        servers.append([f"host-{i}", f"10.0.0.{i}"])

    for i in range(3, 6):
        ws = wb.create_sheet(f"S{i}")
        ws.append(["h"])
        ws.append([i])

    wb.save(path)


def _xlsx_bytes() -> bytes:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    apps = wb.active
    apps.title = "应用系统"
    apps.append(["名称(*)", "类型"])
    apps.append(["系统A", "业务"])
    servers = wb.create_sheet("服务器")
    servers.append(["主机名", "IP"])
    servers.append(["host-0", "10.0.0.0"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_extract_multi_sheet_headers_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    _write_multi_sheet_xlsx(path)

    doc = extract_workbook_tables(str(path), max_sheets=20, max_rows=5000, max_cols=64)
    assert doc["schema_version"] == "spreadsheet_tables@v1"
    assert doc["sheet_count"] == 5
    names = [s["name"] for s in doc["sheets"]]
    assert "应用系统" in names
    assert "服务器" in names
    app = next(s for s in doc["sheets"] if s["name"] == "应用系统")
    assert "名称(*)" in app["headers"]
    assert app["row_count"] >= 2
    assert len(app["rows"][0]) == len(app["headers"])
    server = next(s for s in doc["sheets"] if s["name"] == "服务器")
    assert "IP" in server["headers"]
    assert server["row_count"] >= 10


def test_parser_structured_content_envelope(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    _write_multi_sheet_xlsx(path)

    parser = SpreadsheetTablesParser()
    ctx = ParseContext(
        content_hash="test-hash",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="multi.xlsx",
        size_bytes=path.stat().st_size,
    )
    content = parser.parse(str(path), ctx)
    assert content.source.profile == "tables"
    assert content.metadata.get("format") == "spreadsheet_tables_json"
    assert len(content.tables) == 5
    tj = content.metadata.get("tables_json")
    assert isinstance(tj, dict)
    assert tj.get("sheet_count") == 5


def test_max_sheets_truncation(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "S1"
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    for i in range(2, 6):
        ws = wb.create_sheet(f"S{i}")
        ws.append(["h"])
        ws.append([i])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    doc = extract_workbook_tables(str(path), max_sheets=2, max_rows=100, max_cols=16)
    assert doc["sheet_count"] == 2
    assert any("max_sheets" in n for n in doc.get("notes") or [])


def test_tables_in_supported_profiles() -> None:
    assert "tables" in SUPPORTED_PARSE_PROFILES


def test_upload_auto_parses_tables_content(client, auth_headers) -> None:
    """also_tables 默认 true：上传 xlsx 后 content?profile=tables 可用。"""
    test_client, _ = client
    payload = _xlsx_bytes()
    upload = upload_bytes(
        client,
        auth_headers,
        payload,
        name="multi.xlsx",
        mime=XLSX_MIME,
    )
    assert upload.status_code == 201, upload.text
    reference_id = upload.json()["reference_id"]

    status = test_client.get(f"/v1/files/{reference_id}/status", headers=auth_headers)
    assert status.status_code == 200
    profiles = {p["profile"]: p["status"] for p in status.json()["parse_results"]}
    assert profiles.get("tables") == "completed"

    content = test_client.get(
        f"/v1/files/{reference_id}/content?profile=tables",
        headers=auth_headers,
    )
    assert content.status_code == 200
    body = content.json()
    assert body["ready"] is True
    assert body["source"]["profile"] == "tables"
    tj = body["metadata"]["tables_json"]
    assert tj["schema_version"] == "spreadsheet_tables@v1"
    names = [s["name"] for s in tj["sheets"]]
    assert "应用系统" in names
    assert "服务器" in names


def test_content_tables_sheets_filter(client, auth_headers) -> None:
    test_client, _ = client
    payload = _xlsx_bytes()
    upload = upload_bytes(
        client,
        auth_headers,
        payload,
        name="multi.xlsx",
        mime=XLSX_MIME,
    )
    reference_id = upload.json()["reference_id"]

    content = test_client.get(
        f"/v1/files/{reference_id}/content?profile=tables&sheets=应用系统",
        headers=auth_headers,
    )
    assert content.status_code == 200
    body = content.json()
    assert body["ready"] is True
    tj = body["metadata"]["tables_json"]
    assert tj["sheet_count"] == 1
    assert tj["sheets"][0]["name"] == "应用系统"
    assert body["metadata"].get("sheets_filter")


def test_also_tables_false_skips_auto_tables(client, auth_headers, monkeypatch) -> None:
    reload_parsers(monkeypatch, AUTO_PARSE_SPREADSHEET_ALSO_TABLES="false")
    test_client, _ = client
    payload = _xlsx_bytes()
    upload = upload_bytes(
        client,
        auth_headers,
        payload,
        name="multi.xlsx",
        mime=XLSX_MIME,
    )
    assert upload.status_code == 201
    reference_id = upload.json()["reference_id"]

    status = test_client.get(f"/v1/files/{reference_id}/status", headers=auth_headers)
    profiles = {p["profile"]: p["status"] for p in status.json()["parse_results"]}
    assert "tables" not in profiles or profiles.get("tables") == "not_started"

    content = test_client.get(
        f"/v1/files/{reference_id}/content?profile=tables",
        headers=auth_headers,
    )
    assert content.status_code == 200
    assert content.json()["ready"] is False


def test_trigger_tables_profile(client, auth_headers, monkeypatch) -> None:
    reload_parsers(monkeypatch, AUTO_PARSE_SPREADSHEET_ALSO_TABLES="false")
    test_client, _ = client
    payload = _xlsx_bytes()
    upload = upload_bytes(
        client,
        auth_headers,
        payload,
        name="multi.xlsx",
        mime=XLSX_MIME,
        auto_parse="false",
    )
    reference_id = upload.json()["reference_id"]

    trigger = test_client.post(
        f"/v1/files/{reference_id}/parse",
        headers=auth_headers,
        json={"profile": "tables", "force": False},
    )
    assert trigger.status_code == 202, trigger.text
    assert trigger.json()["profile"] == "tables"
    assert trigger.json()["status"] == "completed"

    content = test_client.get(
        f"/v1/files/{reference_id}/content?profile=tables",
        headers=auth_headers,
    )
    assert content.status_code == 200
    body = content.json()
    assert body["ready"] is True
    assert body["metadata"]["tables_json"]["sheet_count"] >= 1
