"""按文件类型的元数据提取契约测试。"""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.config import Settings
from app.parsers.classifiers import classify_file
from app.parsers.factory import create_processing_factory
from app.parsers.metadata import MetadataParser
from app.parsers.metadata_handlers import (
    AudioMetadataExtractor,
    ImageMetadataExtractor,
    OfficeMetadataExtractor,
    PdfDigitalMetadataExtractor,
    PdfScannedMetadataExtractor,
    SpreadsheetMetadataExtractor,
    TextMetadataExtractor,
    VideoMetadataExtractor,
)
from app.parsers.types import FileClassification, FileKind


def _classify(filename: str, mime: str, data: bytes) -> FileClassification:
    return classify_file(filename, mime, data)


def test_pdf_digital_metadata_page_count_and_text_layer():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    data = buf.getvalue()

    classification = _classify("doc.pdf", "application/pdf", data)
    result = PdfDigitalMetadataExtractor().extract(
        io.BytesIO(data), "doc.pdf", "application/pdf", classification=classification
    )

    assert result.extra.get("page_count") == 1
    assert result.extra.get("text_layer") == "present"
    assert result.extra.get("recommended_profiles") == ["layout"]
    assert "first_page_text_chars" in result.extra
    assert result.extra.get("encrypted") is False


def test_pdf_scanned_metadata_recommends_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    data = buf.getvalue()

    classification = FileClassification(
        kind=FileKind.PDF, subtype="scanned", mime_type="application/pdf"
    )
    result = PdfScannedMetadataExtractor().extract(
        io.BytesIO(data), "scan.pdf", "application/pdf", classification=classification
    )
    assert result.extra["text_layer"] == "absent"
    assert result.extra["ocr_recommended"] is True
    assert "ocr" in result.extra["recommended_profiles"]


def test_spreadsheet_xlsx_metadata_sheet_names():
    wb = Workbook()
    wb.active.title = "概览"
    wb.create_sheet("明细")
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    classification = _classify(
        "book.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data,
    )
    result = SpreadsheetMetadataExtractor().extract(
        io.BytesIO(data),
        "book.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        classification=classification,
    )

    assert result.extra["sheet_count"] == 2
    assert result.extra["sheet_names"] == ["概览", "明细"]
    assert result.extra["workbook_format"] == "xlsx"
    assert len(result.extra["sheets"]) == 2


def test_spreadsheet_xls_metadata_from_sample_if_present():
    sample = Path(__file__).resolve().parents[1] / "sample" / "outbound.xls"
    if not sample.is_file():
        pytest.skip("sample/outbound.xls not present")

    data = sample.read_bytes()
    classification = _classify("outbound.xls", "application/vnd.ms-excel", data)
    assert classification.subtype == "xls"

    result = SpreadsheetMetadataExtractor().extract(
        io.BytesIO(data), "outbound.xls", "application/vnd.ms-excel", classification=classification
    )
    assert result.extra.get("workbook_format") == "xls"
    assert result.extra.get("sheet_count", 0) >= 1
    assert isinstance(result.extra.get("sheet_names"), list)


def test_office_docx_metadata_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_heading("采购说明", level=1)
    document.add_paragraph("这是正文第一段。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    buf = io.BytesIO()
    document.save(buf)
    data = buf.getvalue()

    classification = _classify(
        "spec.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
    )
    assert classification.kind == FileKind.OFFICE

    result = OfficeMetadataExtractor().extract(
        io.BytesIO(data),
        "spec.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        classification=classification,
    )

    assert result.extra["office_format"] == "docx"
    assert result.extra["paragraph_count"] >= 2
    assert result.extra["table_count"] == 1
    assert result.extra["heading_count"] >= 1
    assert any("采购" in h for h in result.extra.get("headings_preview", []))


def test_office_legacy_doc_marks_limited_metadata():
    classification = FileClassification(kind=FileKind.OFFICE, subtype="doc", mime_type="application/msword")
    result = OfficeMetadataExtractor().extract(
        io.BytesIO(b"\xd0\xcf\x11\xe0"),  # OLE magic-ish
        "old.doc",
        "application/msword",
        classification=classification,
    )
    assert result.extra["office_format"] == "doc"
    assert "legacy_doc_limited_metadata" in result.extra.get("extractor_warning", "")


def test_text_metadata_encoding_and_line_count():
    raw = "第一行\n第二行\n".encode("utf-8")
    classification = _classify("notes.txt", "text/plain", raw)
    result = TextMetadataExtractor().extract(
        io.BytesIO(raw), "notes.txt", "text/plain", classification=classification
    )
    assert result.extra["line_count"] == 2
    assert result.extra["char_count"] > 0
    assert result.extra["encoding"] in {"utf-8", "utf-8-sig"}
    assert result.extra["is_empty"] is False


def test_text_csv_metadata_column_hints():
    raw = b"name,qty,price\na,1,2.5\nb,3,4.0\n"
    classification = _classify("items.csv", "text/csv", raw)
    result = TextMetadataExtractor().extract(
        io.BytesIO(raw), "items.csv", "text/csv", classification=classification
    )
    assert result.extra["csv_column_count"] == 3
    assert result.extra["csv_row_count"] == 3
    assert result.extra["csv_dialect"] == ","


def test_text_markdown_heading_preview():
    raw = b"# Title\n\n## Section\n\nbody\n"
    classification = _classify("readme.md", "text/markdown", raw)
    result = TextMetadataExtractor().extract(
        io.BytesIO(raw), "readme.md", "text/markdown", classification=classification
    )
    assert result.extra["markdown_heading_count"] == 2
    assert "Title" in result.extra["headings_preview"][0]


def test_image_metadata_dimensions():
    buf = io.BytesIO()
    Image.new("RGB", (64, 48), color=(10, 20, 30)).save(buf, format="PNG")
    data = buf.getvalue()

    classification = _classify("photo.png", "image/png", data)
    result = ImageMetadataExtractor().extract(
        io.BytesIO(data), "photo.png", "image/png", classification=classification
    )
    assert result.extra["width"] == 64
    assert result.extra["height"] == 48
    assert result.extra["image_format"] == "PNG"
    assert "deep" in result.extra["recommended_profiles"]


def test_factory_registers_office_metadata_extractor():
    factory = create_processing_factory(Settings(layout_parser_name="simple"))
    classification = FileClassification(
        kind=FileKind.OFFICE,
        subtype="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extractor = factory.resolve_metadata_extractor(classification)
    assert extractor.name == "office-metadata"
    routes = factory.build_routes(classification)
    assert any(r.handler_name == "office-metadata" for r in routes)


def test_metadata_parser_injects_processing_routes():
    raw = b"hello\n"
    result = MetadataParser().extract(io.BytesIO(raw), "a.txt", "text/plain")
    assert result.extra["file_kind"] == "text"
    assert result.extra["file_subtype"] == "plain"
    assert isinstance(result.extra.get("processing_routes"), list)
    assert result.extra["processing_routes"][0]["level"] == "metadata"


def test_audio_wav_duration_from_sample():
    path = Path(__file__).resolve().parent.parent / "sample" / "test_silent.wav"
    data = path.read_bytes()
    classification = _classify("test_silent.wav", "audio/wav", data)
    result = AudioMetadataExtractor().extract(
        io.BytesIO(data),
        "test_silent.wav",
        "audio/wav",
        classification=classification,
    )
    assert result.extra.get("duration_seconds") is not None
    assert result.extra.get("sample_rate")
    assert result.extra.get("channels")


def test_video_metadata_does_not_block():
    path = Path(__file__).resolve().parent.parent / "sample" / "test_tiny.mp4"
    data = path.read_bytes()
    classification = _classify("test_tiny.mp4", "video/mp4", data)
    result = VideoMetadataExtractor().extract(
        io.BytesIO(data),
        "test_tiny.mp4",
        "video/mp4",
        classification=classification,
    )
    # duration optional (mutagen may or may not parse); must not raise and keep size
    assert result.size_bytes == len(data)
